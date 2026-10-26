#!/usr/bin/env python3
"""Export the audited TSV as editable XLSX and a publication-ready English PDF."""

import os
import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, A4

OUT = Path(
    os.environ.get(
        "TABLE_S10_OUT_DIR",
        str(Path(__file__).resolve().parents[3] / "results/06_MS/TableS10"),
    )
)
with open(str(OUT / "TableS10.tsv"), newline="") as f:
    reader = csv.reader(f, delimiter="\t")
    data = list(reader)
title = "Table S10"
wb = Workbook()
ws = wb.active
ws.title = "Table S10"
ws.append([title])
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
ws["A1"].font = Font(name="Arial", size=12, bold=True)
ws["A1"].alignment = Alignment(wrap_text=True, vertical="center")
ws.row_dimensions[1].height = 35
ws.append(data[0])
for r in data[1:]:
    ws.append([r[0]] + [int(v) for v in r[1:]])
for col, w in zip("ABCDEFG", [17, 28, 29, 29, 34, 31, 31]):
    ws.column_dimensions[col].width = w
for cell in ws[2]:
    cell.font = Font(name="Arial", size=10, bold=True)
    cell.fill = PatternFill("solid", fgColor="E8EDF2")
    cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    cell.border = Border(top=Side(style="thin"), bottom=Side(style="thin"))
ws.row_dimensions[2].height = 65
for row in ws.iter_rows(min_row=3, max_row=6):
    for c in row:
        c.font = Font(name="Arial", size=11)
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row[0].row].height = 25
ws.freeze_panes = "B3"
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.orientation = "landscape"
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.print_area = "A1:G6"
wb.save(str(OUT / "TableS10.xlsx"))

head = ParagraphStyle(
    "head", fontName="Helvetica-Bold", fontSize=9, leading=12, alignment=1
)
body = ParagraphStyle("body", fontName="Helvetica", fontSize=9, leading=12)
heading = ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=12, leading=16)
grid = [[Paragraph(x, head) for x in data[0]]] + data[1:]
table = Table(
    grid, colWidths=[53, 111, 117, 117, 122, 123, 123], rowHeights=[67, 27, 27, 27, 27]
)
table.setStyle(
    TableStyle(
        [
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 1), (-1, -1), 11),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EDF2")),
            ("LINEABOVE", (0, 0), (-1, 0), 0.8, colors.black),
            ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.black),
            ("LINEBELOW", (0, -1), (-1, -1), 0.8, colors.black),
        ]
    )
)
doc = SimpleDocTemplate(
    str(OUT / "TableS10.pdf"),
    pagesize=landscape(A4),
    leftMargin=27,
    rightMargin=27,
    topMargin=27,
    bottomMargin=27,
)
doc.build([Paragraph(title, heading), Spacer(1, 15), table])
print("Created TableS10.xlsx and TableS10.pdf")
